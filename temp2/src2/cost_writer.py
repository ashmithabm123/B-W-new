import logging
from copy import copy
from datetime import datetime
from openpyxl import load_workbook

log = logging.getLogger("remittance")


def write_cost_sheet(path, sheet_name, cost_df, title, thru_date):
    if thru_date.upper() == "AUTO":
        thru_date = datetime.now().strftime("%m-%d-%Y")

    wb = load_workbook(path)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Output template is missing the '{sheet_name}' sheet")
    ws = wb[sheet_name]

    ws["B1"] = f"{title.upper()} COST REPORT"
    ws["B2"] = f"Costs Thru {thru_date}"
    header_row = 7
    data_start = header_row + 1
    for row_idx in range(data_start, ws.max_row + 1):
        for col_idx in range(1, max(ws.max_column, len(cost_df.columns)) + 1):
            ws.cell(row=row_idx, column=col_idx).value = None

    template_max_row = ws.max_row
    for r, row in enumerate(cost_df.itertuples(index=False), start=data_start):
        if r > template_max_row:
            ws.row_dimensions[r].height = ws.row_dimensions[data_start].height
            for col in range(1, max(ws.max_column, len(cost_df.columns)) + 1):
                ws.cell(r, col)._style = copy(ws.cell(data_start, col)._style)
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(path)
    wb.close()
    log.info("Sheet '%s' written with %d rows", sheet_name, len(cost_df))